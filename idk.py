import pypdf
from docx import Document
def q1by1():
    doc=Document()
    quet=''#to temporarily store a question
    quep=''#to premanently store a question
    reader=pypdf.PdfReader('as2.pdf')
    for page in reader.pages:
        words=page.extract_text(extraction_mode='layout')
        lines=words.split('\n')
        for line in lines:
            que_no=line[0:2]
            match que_no:
                case x if x.endswith('.'):
                    quep=quet
                    quet=''
                    quet+=line
                    print(quep+'!!!!!')
                    doc.add_paragraph(quep)
                    match quep[0:2]:
                        case x if x.endswith('.'):
                            ans=input(f'answer to ques{quep[0:2]}: ')
                            doc.add_paragraph(ans)
                case _:
                    quet+=line+'\n'
    print(quet+'!!!!') #to print out any value that is left in quet at the end 
    doc.add_paragraph(quet)
    match quet[0:2]:
        case x if x.endswith('.'):
            ans=input(f'answer to ques{quet[0:2]}: ')
            doc.add_paragraph(ans)
    doc.save('tosubmit.docx')
q1by1()


#programme can separate out each question
#work on user input answer writing 